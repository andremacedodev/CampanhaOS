"""
Gera o arquivo Excel (.xlsx) do extrato financeiro — detalhe técnico de
apresentação, por isso mora na infraestrutura, não na aplicação. Recebe
o DTO já pronto (FinanceStatementOutput) e só cuida de "como desenhar
isso numa planilha".
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.application.finance.dto import FinanceStatementOutput

_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CURRENCY_FORMAT = '"R$" #,##0.00'

_TYPE_LABELS = {"receita": "Receita", "despesa": "Despesa", "doacao": "Doação"}


def _style_header_row(ws: Worksheet, num_columns: int) -> None:
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit_columns(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def generate_finance_statement_excel(statement: FinanceStatementOutput) -> bytes:
    workbook = Workbook()

    _write_extrato_sheet(workbook.active, statement)

    # Só cria a segunda aba se houver pendência — planilha com aba vazia
    # é mais confuso que simplesmente não ter a aba.
    if statement.pending_expenses:
        pending_sheet = workbook.create_sheet("Pendências")
        _write_pendencias_sheet(pending_sheet, statement)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_extrato_sheet(ws: Worksheet, statement: FinanceStatementOutput) -> None:
    ws.title = "Extrato"
    ws.append(["Data", "Tipo", "Categoria", "Descrição", "Valor", "Saldo Acumulado"])
    _style_header_row(ws, 6)

    for line in statement.lines:
        row = [
            line.date.strftime("%d/%m/%Y"),
            _TYPE_LABELS.get(line.type, line.type),
            line.category,
            line.description or "",
            float(line.amount),
            float(line.running_balance),
        ]
        ws.append(row)
        last_row = ws.max_row
        ws.cell(row=last_row, column=5).number_format = _CURRENCY_FORMAT
        ws.cell(row=last_row, column=6).number_format = _CURRENCY_FORMAT
        # Despesa em vermelho — facilita achar visualmente na hora de
        # conferir contra o extrato do banco.
        if line.type == "despesa":
            ws.cell(row=last_row, column=5).font = Font(color="DC2626")

    _autofit_columns(ws, [12, 12, 22, 35, 15, 17])
    ws.freeze_panes = "A2"  # cabeçalho fixo ao rolar


def _write_pendencias_sheet(ws: Worksheet, statement: FinanceStatementOutput) -> None:
    ws.append(["Data do Lançamento", "Categoria", "Descrição", "Valor Total", "Já Pago", "Falta Pagar"])
    _style_header_row(ws, 6)

    for pending in statement.pending_expenses:
        row = [
            pending.occurred_at.strftime("%d/%m/%Y"),
            pending.category,
            pending.description or "",
            float(pending.amount),
            float(pending.amount_paid),
            float(pending.amount_remaining),
        ]
        ws.append(row)
        last_row = ws.max_row
        for col in (4, 5, 6):
            ws.cell(row=last_row, column=col).number_format = _CURRENCY_FORMAT

    _autofit_columns(ws, [16, 22, 35, 15, 15, 15])
    ws.freeze_panes = "A2"
